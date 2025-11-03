"""
Data Migration Script
Migrates data from Excel file to PostgreSQL database
"""

import psycopg2
from openpyxl import load_workbook
from datetime import datetime
from decimal import Decimal
import os
import sys

# Database configuration
DB_CONFIG = {
    "host": os.getenv("DB_HOST", "localhost"),
    "port": os.getenv("DB_PORT", "5432"),
    "database": os.getenv("DB_NAME", "sales_data_hub"),
    "user": os.getenv("DB_USER", "sales_admin"),
    "password": os.getenv("DB_PASSWORD", "sales_secure_pass_2025")
}

EXCEL_FILE = "../Neelam's Data Sheet (2).xlsx"

def get_db_connection():
    """Create database connection"""
    return psycopg2.connect(**DB_CONFIG)

def safe_str(value):
    """Safely convert to string"""
    if value is None:
        return None
    return str(value).strip() if str(value).strip() else None

def safe_decimal(value):
    """Safely convert to Decimal"""
    if value is None or value == "":
        return Decimal("0")
    try:
        if isinstance(value, str):
            # Remove any non-numeric characters except . and -
            value = ''.join(c for c in value if c.isdigit() or c in '.-')
        return Decimal(str(value))
    except:
        return Decimal("0")

def safe_date(value):
    """Safely convert to date"""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except:
            return None
    return None

def safe_bool(value):
    """Safely convert to boolean"""
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.lower() in ['yes', 'true', '1', 'y']
    return bool(value)

def migrate_salespeople(wb, conn):
    """Migrate Salesperson Master data"""
    print("\n📊 Migrating Salespeople...")
    ws = wb['Salesperson Master']
    cur = conn.cursor()

    # Skip header rows (rows 1-4)
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:  # Skip empty rows
            continue

        try:
            cur.execute(
                """INSERT INTO salespeople
                (salesperson_id, name, role, email, phone, joining_date, department, employment_status)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (salesperson_id) DO UPDATE SET
                name = EXCLUDED.name,
                role = EXCLUDED.role,
                email = EXCLUDED.email,
                phone = EXCLUDED.phone,
                joining_date = EXCLUDED.joining_date,
                department = EXCLUDED.department,
                employment_status = EXCLUDED.employment_status""",
                (
                    safe_str(row[0]),  # salesperson_id
                    safe_str(row[1]),  # name
                    safe_str(row[2]),  # role
                    safe_str(row[3]),  # email
                    safe_str(row[4]),  # phone
                    safe_date(row[5]),  # joining_date
                    safe_str(row[6]),  # department
                    safe_str(row[9]) if row[9] else 'Active'  # employment_status
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating salesperson {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} salespeople")

def migrate_customers(wb, conn):
    """Migrate Customer Master data"""
    print("\n👥 Migrating Customers...")
    ws = wb['Customer Master']
    cur = conn.cursor()

    # Skip header rows (rows 1-4)
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:  # Skip empty rows
            continue

        try:
            cur.execute(
                """INSERT INTO customers
                (customer_id, client_name, industry, type, focus_account, cin_number,
                 status, onboarding_date, hunter_id, farmer_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (customer_id) DO UPDATE SET
                client_name = EXCLUDED.client_name,
                industry = EXCLUDED.industry,
                type = EXCLUDED.type,
                focus_account = EXCLUDED.focus_account,
                cin_number = EXCLUDED.cin_number,
                status = EXCLUDED.status,
                onboarding_date = EXCLUDED.onboarding_date,
                hunter_id = EXCLUDED.hunter_id,
                farmer_id = EXCLUDED.farmer_id""",
                (
                    safe_str(row[0]),  # customer_id
                    safe_str(row[1]),  # client_name
                    safe_str(row[2]),  # industry
                    safe_str(row[3]),  # type
                    safe_bool(row[4]),  # focus_account
                    safe_str(row[5]),  # cin_number
                    safe_str(row[6]) if row[6] else 'Active',  # status
                    safe_date(row[7]),  # onboarding_date
                    safe_str(row[8]),  # hunter_id
                    safe_str(row[9])   # farmer_id
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating customer {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} customers")

def migrate_customer_targets(wb, conn):
    """Migrate Customer Targets data"""
    print("\n🎯 Migrating Customer Targets...")
    ws = wb['Customer Master']
    cur = conn.cursor()

    fiscal_year = 2025  # Default fiscal year
    count = 0

    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue

        try:
            # Extract annual targets from columns K and L
            annual_netnew = safe_decimal(row[10])  # Column K
            annual_retention = safe_decimal(row[11])  # Column L

            cur.execute(
                """INSERT INTO customer_targets
                (customer_id, fiscal_year, annual_netnew_target, annual_retention_target,
                 q1_netnew_target, q1_retention_target, q2_netnew_target, q2_retention_target,
                 q3_netnew_target, q3_retention_target, q4_netnew_target, q4_retention_target)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (customer_id, fiscal_year) DO UPDATE SET
                annual_netnew_target = EXCLUDED.annual_netnew_target,
                annual_retention_target = EXCLUDED.annual_retention_target,
                q1_netnew_target = EXCLUDED.q1_netnew_target,
                q1_retention_target = EXCLUDED.q1_retention_target,
                q2_netnew_target = EXCLUDED.q2_netnew_target,
                q2_retention_target = EXCLUDED.q2_retention_target,
                q3_netnew_target = EXCLUDED.q3_netnew_target,
                q3_retention_target = EXCLUDED.q3_retention_target,
                q4_netnew_target = EXCLUDED.q4_netnew_target,
                q4_retention_target = EXCLUDED.q4_retention_target""",
                (
                    safe_str(row[0]),  # customer_id
                    fiscal_year,
                    annual_netnew,
                    annual_retention,
                    annual_netnew / 4,  # Q1 netnew
                    annual_retention / 4,  # Q1 retention
                    annual_netnew / 4,  # Q2 netnew
                    annual_retention / 4,  # Q2 retention
                    annual_netnew / 4,  # Q3 netnew
                    annual_retention / 4,  # Q3 retention
                    annual_netnew / 4,  # Q4 netnew
                    annual_retention / 4   # Q4 retention
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating customer target {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} customer targets")

def migrate_salesperson_targets(wb, conn):
    """Migrate Salesperson Targets data"""
    print("\n🎯 Migrating Salesperson Targets...")
    ws = wb['Salesperson Master']
    cur = conn.cursor()

    fiscal_year = 2025
    count = 0

    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue

        try:
            annual_netnew = safe_decimal(row[10])  # Column K
            annual_retention = safe_decimal(row[11])  # Column L

            cur.execute(
                """INSERT INTO salesperson_targets
                (salesperson_id, fiscal_year, annual_netnew_target, annual_retention_target,
                 q1_netnew_target, q1_retention_target, q2_netnew_target, q2_retention_target,
                 q3_netnew_target, q3_retention_target, q4_netnew_target, q4_retention_target)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (salesperson_id, fiscal_year) DO UPDATE SET
                annual_netnew_target = EXCLUDED.annual_netnew_target,
                annual_retention_target = EXCLUDED.annual_retention_target,
                q1_netnew_target = EXCLUDED.q1_netnew_target,
                q1_retention_target = EXCLUDED.q1_retention_target,
                q2_netnew_target = EXCLUDED.q2_netnew_target,
                q2_retention_target = EXCLUDED.q2_retention_target,
                q3_netnew_target = EXCLUDED.q3_netnew_target,
                q3_retention_target = EXCLUDED.q3_retention_target,
                q4_netnew_target = EXCLUDED.q4_netnew_target,
                q4_retention_target = EXCLUDED.q4_retention_target""",
                (
                    safe_str(row[0]),
                    fiscal_year,
                    annual_netnew,
                    annual_retention,
                    annual_netnew / 4,
                    annual_retention / 4,
                    annual_netnew / 4,
                    annual_retention / 4,
                    annual_netnew / 4,
                    annual_retention / 4,
                    annual_netnew / 4,
                    annual_retention / 4
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating salesperson target {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} salesperson targets")

def migrate_products(wb, conn):
    """Migrate Product Master data"""
    print("\n📦 Migrating Products...")
    ws = wb['Product Master']
    cur = conn.cursor()

    count = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue

        try:
            cur.execute(
                """INSERT INTO products
                (product_id, product_name, category, base_price, description, is_active)
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (product_id) DO UPDATE SET
                product_name = EXCLUDED.product_name,
                category = EXCLUDED.category,
                base_price = EXCLUDED.base_price,
                description = EXCLUDED.description,
                is_active = EXCLUDED.is_active""",
                (
                    safe_str(row[0]),  # product_id
                    safe_str(row[1]),  # product_name
                    safe_str(row[2]),  # category
                    safe_decimal(row[3]),  # base_price
                    safe_str(row[4]),  # description
                    True
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating product {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} products")

def migrate_invoices(wb, conn):
    """Migrate Invoice Transactions data"""
    print("\n💰 Migrating Invoices...")
    ws = wb['Invoice Transactions']
    cur = conn.cursor()

    count = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:  # Skip empty rows
            continue

        try:
            # Determine fiscal year from invoice date
            invoice_date = safe_date(row[3])
            fiscal_year = invoice_date.year if invoice_date else 2025

            cur.execute(
                """INSERT INTO invoices
                (invoice_id, customer_id, invoice_date, usage_month, account_type,
                 revenue_type, product_service, revenue_category, quarter, amount, gst,
                 billing_status, owner_id, notes, fiscal_year)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (invoice_id) DO UPDATE SET
                customer_id = EXCLUDED.customer_id,
                invoice_date = EXCLUDED.invoice_date,
                usage_month = EXCLUDED.usage_month,
                account_type = EXCLUDED.account_type,
                revenue_type = EXCLUDED.revenue_type,
                product_service = EXCLUDED.product_service,
                revenue_category = EXCLUDED.revenue_category,
                quarter = EXCLUDED.quarter,
                amount = EXCLUDED.amount,
                gst = EXCLUDED.gst,
                billing_status = EXCLUDED.billing_status,
                owner_id = EXCLUDED.owner_id,
                notes = EXCLUDED.notes,
                fiscal_year = EXCLUDED.fiscal_year""",
                (
                    safe_str(row[0]),  # invoice_id
                    safe_str(row[1]),  # customer_id
                    invoice_date,  # invoice_date
                    safe_date(row[4]),  # usage_month
                    safe_str(row[5]),  # account_type
                    safe_str(row[6]),  # revenue_type
                    safe_str(row[7]),  # product_service
                    safe_str(row[8]),  # revenue_category (Net New/Retention)
                    safe_str(row[9]),  # quarter
                    safe_decimal(row[11]),  # amount
                    safe_decimal(row[12]),  # gst
                    safe_str(row[14]) if row[14] else 'Unbilled',  # billing_status
                    safe_str(row[15]),  # owner_id
                    safe_str(row[16]),  # notes
                    fiscal_year
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating invoice {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} invoices")

def migrate_agreements(wb, conn):
    """Migrate Agreements data"""
    print("\n📄 Migrating Agreements...")
    ws = wb['Agreements']
    cur = conn.cursor()

    count = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue

        try:
            cur.execute(
                """INSERT INTO agreements
                (agreement_id, customer_id, agreement_type, start_date, end_date,
                 contract_value, payment_terms, auto_renewal, notice_period_days, status, notes)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (agreement_id) DO UPDATE SET
                customer_id = EXCLUDED.customer_id,
                agreement_type = EXCLUDED.agreement_type,
                start_date = EXCLUDED.start_date,
                end_date = EXCLUDED.end_date,
                contract_value = EXCLUDED.contract_value,
                payment_terms = EXCLUDED.payment_terms,
                auto_renewal = EXCLUDED.auto_renewal,
                notice_period_days = EXCLUDED.notice_period_days,
                status = EXCLUDED.status,
                notes = EXCLUDED.notes""",
                (
                    safe_str(row[0]),  # agreement_id
                    safe_str(row[1]),  # customer_id
                    safe_str(row[3]),  # agreement_type
                    safe_date(row[4]),  # start_date
                    safe_date(row[5]),  # end_date
                    safe_decimal(row[6]),  # contract_value
                    safe_str(row[7]),  # payment_terms
                    safe_bool(row[8]),  # auto_renewal
                    int(row[9]) if row[9] else None,  # notice_period_days
                    safe_str(row[10]) if row[10] else 'Active',  # status
                    safe_str(row[11])  # notes
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating agreement {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} agreements")

def migrate_deals(wb, conn):
    """Migrate HubSpot Pipeline data"""
    print("\n🤝 Migrating Deals...")
    ws = wb['HubSpot Pipeline']
    cur = conn.cursor()

    count = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue

        try:
            close_date = safe_date(row[9])
            fiscal_year = close_date.year if close_date else 2025

            cur.execute(
                """INSERT INTO deals
                (deal_id, deal_name, customer_id, owner_id, stage, amount, arr,
                 one_time_amount, expected_golive_date, close_date, q1_revenue,
                 q2_revenue, q3_revenue, q4_revenue, fiscal_year)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (deal_id) DO UPDATE SET
                deal_name = EXCLUDED.deal_name,
                customer_id = EXCLUDED.customer_id,
                owner_id = EXCLUDED.owner_id,
                stage = EXCLUDED.stage,
                amount = EXCLUDED.amount,
                arr = EXCLUDED.arr,
                one_time_amount = EXCLUDED.one_time_amount,
                expected_golive_date = EXCLUDED.expected_golive_date,
                close_date = EXCLUDED.close_date,
                q1_revenue = EXCLUDED.q1_revenue,
                q2_revenue = EXCLUDED.q2_revenue,
                q3_revenue = EXCLUDED.q3_revenue,
                q4_revenue = EXCLUDED.q4_revenue,
                fiscal_year = EXCLUDED.fiscal_year""",
                (
                    safe_str(row[0]),  # deal_id
                    safe_str(row[1]),  # deal_name
                    safe_str(row[2]),  # customer_id
                    safe_str(row[3]),  # owner_id
                    safe_str(row[4]),  # stage
                    safe_decimal(row[5]),  # amount
                    safe_decimal(row[6]),  # arr
                    safe_decimal(row[7]),  # one_time_amount
                    safe_date(row[8]),  # expected_golive_date
                    close_date,  # close_date
                    safe_decimal(row[10]),  # q1_revenue
                    safe_decimal(row[11]),  # q2_revenue
                    safe_decimal(row[12]),  # q3_revenue
                    safe_decimal(row[13]),  # q4_revenue
                    fiscal_year
                )
            )
            count += 1
        except Exception as e:
            print(f"  ⚠ Error migrating deal {row[0]}: {str(e)}")
            continue

    conn.commit()
    print(f"  ✓ Migrated {count} deals")

def create_collections_from_invoices(conn):
    """Create collection records from invoices"""
    print("\n💳 Creating Collections from Invoices...")
    cur = conn.cursor()

    cur.execute("CALL sp_create_collections_from_invoices()")
    conn.commit()

    cur.execute("SELECT COUNT(*) FROM collections")
    count = cur.fetchone()[0]
    print(f"  ✓ Created {count} collection records")

def main():
    """Main migration function"""
    print("=" * 60)
    print("Sales Data Management Hub - Data Migration")
    print("=" * 60)

    # Check if Excel file exists
    if not os.path.exists(EXCEL_FILE):
        print(f"\n❌ Error: Excel file not found at {EXCEL_FILE}")
        sys.exit(1)

    print(f"\n📂 Loading Excel file: {EXCEL_FILE}")
    try:
        wb = load_workbook(EXCEL_FILE)
        print(f"  ✓ Excel file loaded successfully")
        print(f"  ✓ Found {len(wb.sheetnames)} sheets")
    except Exception as e:
        print(f"\n❌ Error loading Excel file: {str(e)}")
        sys.exit(1)

    # Connect to database
    print(f"\n🔌 Connecting to database...")
    try:
        conn = get_db_connection()
        print(f"  ✓ Connected to {DB_CONFIG['database']}")
    except Exception as e:
        print(f"\n❌ Error connecting to database: {str(e)}")
        sys.exit(1)

    try:
        # Run migrations in order (respecting foreign key constraints)
        migrate_salespeople(wb, conn)
        migrate_customers(wb, conn)
        migrate_products(wb, conn)
        migrate_customer_targets(wb, conn)
        migrate_salesperson_targets(wb, conn)
        migrate_invoices(wb, conn)
        migrate_agreements(wb, conn)
        migrate_deals(wb, conn)
        create_collections_from_invoices(conn)

        print("\n" + "=" * 60)
        print("✅ Migration completed successfully!")
        print("=" * 60)
        print("\nYou can now:")
        print("  1. Query the database using SQL")
        print("  2. Start the API server: python3 backend/app.py")
        print("  3. Access the API at: http://localhost:8000")
        print("")

    except Exception as e:
        print(f"\n❌ Migration failed: {str(e)}")
        sys.exit(1)
    finally:
        conn.close()

if __name__ == "__main__":
    main()
