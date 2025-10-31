"""
Quick Demo Viewer for Sales Data Management Hub
Reads Excel file and displays data in a web interface
"""

from fastapi import FastAPI
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from datetime import datetime
from decimal import Decimal
import uvicorn

app = FastAPI(title="Sales Data Management Hub - Demo Viewer")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

EXCEL_FILE = "Neelam's Data Sheet (2).xlsx"

def safe_str(value):
    """Safely convert to string"""
    if value is None:
        return ""
    return str(value).strip()

def safe_number(value):
    """Safely convert to number"""
    if value is None or value == "":
        return 0
    try:
        if isinstance(value, str) and value.startswith('='):
            return 0  # Formula
        return float(value)
    except:
        return 0

def load_excel_data():
    """Load data from Excel file"""
    wb = load_workbook(EXCEL_FILE)
    return wb

@app.get("/", response_class=HTMLResponse)
def read_root():
    """Home page with navigation"""
    return """
    <!DOCTYPE html>
    <html>
    <head>
        <title>Sales Data Management Hub - Demo</title>
        <style>
            * { margin: 0; padding: 0; box-sizing: border-box; }
            body {
                font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                min-height: 100vh;
                padding: 20px;
            }
            .container {
                max-width: 1200px;
                margin: 0 auto;
            }
            .header {
                background: white;
                border-radius: 15px;
                padding: 40px;
                margin-bottom: 30px;
                box-shadow: 0 10px 30px rgba(0,0,0,0.2);
                text-align: center;
            }
            h1 {
                color: #667eea;
                font-size: 2.5em;
                margin-bottom: 10px;
            }
            .subtitle {
                color: #666;
                font-size: 1.2em;
            }
            .cards {
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(300px, 1fr));
                gap: 20px;
                margin-bottom: 30px;
            }
            .card {
                background: white;
                border-radius: 15px;
                padding: 30px;
                box-shadow: 0 5px 20px rgba(0,0,0,0.1);
                transition: transform 0.3s ease, box-shadow 0.3s ease;
                cursor: pointer;
            }
            .card:hover {
                transform: translateY(-5px);
                box-shadow: 0 10px 30px rgba(0,0,0,0.2);
            }
            .card h2 {
                color: #667eea;
                margin-bottom: 15px;
                font-size: 1.5em;
            }
            .card p {
                color: #666;
                line-height: 1.6;
            }
            .card .icon {
                font-size: 3em;
                margin-bottom: 15px;
            }
            .btn {
                display: inline-block;
                background: #667eea;
                color: white;
                padding: 12px 30px;
                border-radius: 8px;
                text-decoration: none;
                margin-top: 15px;
                transition: background 0.3s ease;
                font-weight: 600;
            }
            .btn:hover {
                background: #5568d3;
            }
            .features {
                background: white;
                border-radius: 15px;
                padding: 40px;
                box-shadow: 0 5px 20px rgba(0,0,0,0.1);
            }
            .features h3 {
                color: #667eea;
                margin-bottom: 20px;
                font-size: 1.8em;
            }
            .features ul {
                list-style: none;
                padding-left: 0;
            }
            .features li {
                padding: 12px 0;
                border-bottom: 1px solid #eee;
                color: #555;
            }
            .features li:before {
                content: "✓ ";
                color: #667eea;
                font-weight: bold;
                margin-right: 10px;
            }
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <h1>📊 Sales Data Management Hub</h1>
                <p class="subtitle">Dynamic Data Management System - Demo Viewer</p>
            </div>

            <div class="cards">
                <div class="card" onclick="window.location.href='/dashboard'">
                    <div class="icon">📈</div>
                    <h2>Dashboard</h2>
                    <p>View key metrics, team performance, and executive summary</p>
                    <a href="/dashboard" class="btn">View Dashboard</a>
                </div>

                <div class="card" onclick="window.location.href='/customers'">
                    <div class="icon">👥</div>
                    <h2>Customers</h2>
                    <p>Browse all customers with targets and achievements</p>
                    <a href="/customers" class="btn">View Customers</a>
                </div>

                <div class="card" onclick="window.location.href='/salespeople'">
                    <div class="icon">👤</div>
                    <h2>Salespeople</h2>
                    <p>View sales team performance and quarterly achievements</p>
                    <a href="/salespeople" class="btn">View Team</a>
                </div>

                <div class="card" onclick="window.location.href='/invoices'">
                    <div class="icon">💰</div>
                    <h2>Invoices</h2>
                    <p>Browse all invoice transactions and revenue data</p>
                    <a href="/invoices" class="btn">View Invoices</a>
                </div>

                <div class="card" onclick="window.location.href='/collections'">
                    <div class="icon">💳</div>
                    <h2>Collections</h2>
                    <p>Track payments, outstanding amounts, and efficiency</p>
                    <a href="/collections" class="btn">View Collections</a>
                </div>

                <div class="card" onclick="window.location.href='/api-docs'">
                    <div class="icon">📚</div>
                    <h2>API Docs</h2>
                    <p>Interactive API documentation and testing</p>
                    <a href="/docs" class="btn">View API Docs</a>
                </div>
            </div>

            <div class="features">
                <h3>🎯 System Features</h3>
                <ul>
                    <li>Complete PostgreSQL database with 9 interconnected tables</li>
                    <li>RESTful API with 20+ endpoints for all operations</li>
                    <li>Automated business logic: quarterly targets, achievements, gap analysis</li>
                    <li>Hunter/Farmer model with Net New vs Retention tracking</li>
                    <li>Real-time dashboard metrics and analytics views</li>
                    <li>Auto-calculations for GST, totals, quarters, fiscal years</li>
                    <li>Collections tracking with payment efficiency monitoring</li>
                    <li>Data migration from Excel to SQL with full preservation</li>
                </ul>
            </div>
        </div>
    </body>
    </html>
    """

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard():
    """Dashboard view"""
    try:
        wb = load_excel_data()
        ws = wb['📊 Dashboard']

        # Extract key metrics from dashboard
        total_customers = safe_str(ws['D5'].value) if 'D5' in ws else "Loading..."
        total_salespeople = safe_str(ws['D6'].value) if 'D6' in ws else "Loading..."
        total_products = safe_str(ws['D7'].value) if 'D7' in ws else "Loading..."
        active_deals = safe_str(ws['D8'].value) if 'D8' in ws else "Loading..."

        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Dashboard - Sales Data Hub</title>
            <style>
                body {{
                    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                    background: #f5f7fa;
                    padding: 20px;
                    margin: 0;
                }}
                .container {{ max-width: 1400px; margin: 0 auto; }}
                .header {{
                    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                    color: white;
                    padding: 30px;
                    border-radius: 15px;
                    margin-bottom: 30px;
                }}
                .metrics {{
                    display: grid;
                    grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
                    gap: 20px;
                    margin-bottom: 30px;
                }}
                .metric-card {{
                    background: white;
                    padding: 25px;
                    border-radius: 12px;
                    box-shadow: 0 2px 10px rgba(0,0,0,0.1);
                }}
                .metric-card h3 {{
                    color: #666;
                    font-size: 0.9em;
                    margin-bottom: 10px;
                    text-transform: uppercase;
                }}
                .metric-card .value {{
                    font-size: 2.5em;
                    color: #667eea;
                    font-weight: bold;
                }}
                .back-btn {{
                    display: inline-block;
                    background: white;
                    color: #667eea;
                    padding: 10px 20px;
                    border-radius: 8px;
                    text-decoration: none;
                    font-weight: 600;
                }}
                .back-btn:hover {{ background: #f0f0f0; }}
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <a href="/" class="back-btn">← Back to Home</a>
                    <h1 style="margin-top: 20px;">📊 Executive Dashboard</h1>
                    <p>Key Performance Indicators</p>
                </div>

                <div class="metrics">
                    <div class="metric-card">
                        <h3>Total Customers</h3>
                        <div class="value">{total_customers}</div>
                    </div>
                    <div class="metric-card">
                        <h3>Total Salespeople</h3>
                        <div class="value">{total_salespeople}</div>
                    </div>
                    <div class="metric-card">
                        <h3>Total Products</h3>
                        <div class="value">{total_products}</div>
                    </div>
                    <div class="metric-card">
                        <h3>Active Deals</h3>
                        <div class="value">{active_deals}</div>
                    </div>
                </div>

                <div class="metric-card">
                    <h3>💡 Dashboard Features in Full System</h3>
                    <ul>
                        <li>Q2 Team Target vs Achievement</li>
                        <li>Total Compensation Budget</li>
                        <li>Average OTE Achievement</li>
                        <li>YTD Revenue (Net New vs Retention)</li>
                        <li>Collection Efficiency</li>
                        <li>Pipeline Value</li>
                    </ul>
                </div>
            </div>
        </body>
        </html>
        """
    except Exception as e:
        return f"<html><body><h1>Error loading dashboard</h1><p>{str(e)}</p></body></html>"

@app.get("/customers", response_class=HTMLResponse)
def customers():
    """Customers view"""
    try:
        wb = load_excel_data()
        ws = wb['Customer Master']

        html = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>Customers - Sales Data Hub</title>
            <style>
                body {
                    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                    background: #f5f7fa;
                    padding: 20px;
                    margin: 0;
                }
                .container { max-width: 1400px; margin: 0 auto; }
                .header {
                    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                    color: white;
                    padding: 30px;
                    border-radius: 15px;
                    margin-bottom: 30px;
                }
                table {
                    width: 100%;
                    background: white;
                    border-radius: 12px;
                    overflow: hidden;
                    box-shadow: 0 2px 10px rgba(0,0,0,0.1);
                    border-collapse: collapse;
                }
                th {
                    background: #667eea;
                    color: white;
                    padding: 15px;
                    text-align: left;
                    font-weight: 600;
                }
                td {
                    padding: 12px 15px;
                    border-bottom: 1px solid #eee;
                }
                tr:hover {
                    background: #f9f9f9;
                }
                .back-btn {
                    display: inline-block;
                    background: white;
                    color: #667eea;
                    padding: 10px 20px;
                    border-radius: 8px;
                    text-decoration: none;
                    font-weight: 600;
                }
                .back-btn:hover { background: #f0f0f0; }
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <a href="/" class="back-btn">← Back to Home</a>
                    <h1 style="margin-top: 20px;">👥 Customer Master</h1>
                    <p>All customers with targets and achievements</p>
                </div>

                <table>
                    <thead>
                        <tr>
                            <th>Customer ID</th>
                            <th>Client Name</th>
                            <th>Industry</th>
                            <th>Type</th>
                            <th>Status</th>
                            <th>Hunter ID</th>
                            <th>Farmer ID</th>
                        </tr>
                    </thead>
                    <tbody>
        """

        # Read customer data (skip header rows)
        count = 0
        for row in ws.iter_rows(min_row=5, max_row=30, values_only=True):
            if not row[0] or count >= 20:  # Limit to first 20 for demo
                break
            html += f"""
                <tr>
                    <td>{safe_str(row[0])}</td>
                    <td><strong>{safe_str(row[1])}</strong></td>
                    <td>{safe_str(row[2])}</td>
                    <td>{safe_str(row[3])}</td>
                    <td>{safe_str(row[6])}</td>
                    <td>{safe_str(row[8])}</td>
                    <td>{safe_str(row[9])}</td>
                </tr>
            """
            count += 1

        html += """
                    </tbody>
                </table>
                <p style="margin-top: 20px; text-align: center; color: #666;">
                    Showing first 20 customers. Full system shows all 1100+ customers with quarterly targets and achievements.
                </p>
            </div>
        </body>
        </html>
        """
        return html
    except Exception as e:
        return f"<html><body><h1>Error loading customers</h1><p>{str(e)}</p></body></html>"

@app.get("/salespeople", response_class=HTMLResponse)
def salespeople():
    """Salespeople view"""
    try:
        wb = load_excel_data()
        ws = wb['Salesperson Master']

        html = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>Salespeople - Sales Data Hub</title>
            <style>
                body {
                    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                    background: #f5f7fa;
                    padding: 20px;
                    margin: 0;
                }
                .container { max-width: 1400px; margin: 0 auto; }
                .header {
                    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                    color: white;
                    padding: 30px;
                    border-radius: 15px;
                    margin-bottom: 30px;
                }
                table {
                    width: 100%;
                    background: white;
                    border-radius: 12px;
                    overflow: hidden;
                    box-shadow: 0 2px 10px rgba(0,0,0,0.1);
                    border-collapse: collapse;
                }
                th {
                    background: #667eea;
                    color: white;
                    padding: 15px;
                    text-align: left;
                    font-weight: 600;
                }
                td {
                    padding: 12px 15px;
                    border-bottom: 1px solid #eee;
                }
                tr:hover {
                    background: #f9f9f9;
                }
                .back-btn {
                    display: inline-block;
                    background: white;
                    color: #667eea;
                    padding: 10px 20px;
                    border-radius: 8px;
                    text-decoration: none;
                    font-weight: 600;
                }
                .back-btn:hover { background: #f0f0f0; }
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <a href="/" class="back-btn">← Back to Home</a>
                    <h1 style="margin-top: 20px;">👤 Salesperson Master</h1>
                    <p>Sales team with roles and targets</p>
                </div>

                <table>
                    <thead>
                        <tr>
                            <th>ID</th>
                            <th>Name</th>
                            <th>Role</th>
                            <th>Email</th>
                            <th>Department</th>
                            <th>Status</th>
                        </tr>
                    </thead>
                    <tbody>
        """

        # Read salesperson data
        for row in ws.iter_rows(min_row=5, max_row=20, values_only=True):
            if not row[0]:
                break
            html += f"""
                <tr>
                    <td>{safe_str(row[0])}</td>
                    <td><strong>{safe_str(row[1])}</strong></td>
                    <td>{safe_str(row[2])}</td>
                    <td>{safe_str(row[3])}</td>
                    <td>{safe_str(row[6])}</td>
                    <td>{safe_str(row[9])}</td>
                </tr>
            """

        html += """
                    </tbody>
                </table>
            </div>
        </body>
        </html>
        """
        return html
    except Exception as e:
        return f"<html><body><h1>Error loading salespeople</h1><p>{str(e)}</p></body></html>"

@app.get("/invoices", response_class=HTMLResponse)
def invoices():
    """Invoices view"""
    try:
        wb = load_excel_data()
        ws = wb['Invoice Transactions']

        html = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>Invoices - Sales Data Hub</title>
            <style>
                body {
                    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                    background: #f5f7fa;
                    padding: 20px;
                    margin: 0;
                }
                .container { max-width: 1400px; margin: 0 auto; }
                .header {
                    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                    color: white;
                    padding: 30px;
                    border-radius: 15px;
                    margin-bottom: 30px;
                }
                table {
                    width: 100%;
                    background: white;
                    border-radius: 12px;
                    overflow: hidden;
                    box-shadow: 0 2px 10px rgba(0,0,0,0.1);
                    border-collapse: collapse;
                    font-size: 0.9em;
                }
                th {
                    background: #667eea;
                    color: white;
                    padding: 12px 10px;
                    text-align: left;
                    font-weight: 600;
                }
                td {
                    padding: 10px;
                    border-bottom: 1px solid #eee;
                }
                tr:hover {
                    background: #f9f9f9;
                }
                .back-btn {
                    display: inline-block;
                    background: white;
                    color: #667eea;
                    padding: 10px 20px;
                    border-radius: 8px;
                    text-decoration: none;
                    font-weight: 600;
                }
                .back-btn:hover { background: #f0f0f0; }
                .badge {
                    display: inline-block;
                    padding: 4px 8px;
                    border-radius: 4px;
                    font-size: 0.85em;
                    font-weight: 600;
                }
                .badge-new { background: #d4edda; color: #155724; }
                .badge-retention { background: #cce5ff; color: #004085; }
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <a href="/" class="back-btn">← Back to Home</a>
                    <h1 style="margin-top: 20px;">💰 Invoice Transactions</h1>
                    <p>All revenue transactions with categorization</p>
                </div>

                <table>
                    <thead>
                        <tr>
                            <th>Invoice ID</th>
                            <th>Customer</th>
                            <th>Date</th>
                            <th>Quarter</th>
                            <th>Category</th>
                            <th>Amount</th>
                            <th>Status</th>
                        </tr>
                    </thead>
                    <tbody>
        """

        # Read invoice data
        count = 0
        for row in ws.iter_rows(min_row=4, max_row=50, values_only=True):
            if not row[0] or count >= 30:
                break

            category = safe_str(row[8])
            badge_class = "badge-new" if category == "Net New" else "badge-retention"

            html += f"""
                <tr>
                    <td>{safe_str(row[0])}</td>
                    <td>{safe_str(row[1])}</td>
                    <td>{safe_str(row[3])[:10] if row[3] else ''}</td>
                    <td><strong>{safe_str(row[9])}</strong></td>
                    <td><span class="badge {badge_class}">{category}</span></td>
                    <td>₹ {safe_number(row[11]):,.0f}</td>
                    <td>{safe_str(row[14])}</td>
                </tr>
            """
            count += 1

        html += """
                    </tbody>
                </table>
                <p style="margin-top: 20px; text-align: center; color: #666;">
                    Showing first 30 invoices. Full system has 350+ transactions with complete tracking.
                </p>
            </div>
        </body>
        </html>
        """
        return html
    except Exception as e:
        return f"<html><body><h1>Error loading invoices</h1><p>{str(e)}</p></body></html>"

@app.get("/collections", response_class=HTMLResponse)
def collections():
    """Collections view"""
    return """
    <!DOCTYPE html>
    <html>
    <head>
        <title>Collections - Sales Data Hub</title>
        <style>
            body {
                font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                background: #f5f7fa;
                padding: 20px;
                margin: 0;
            }
            .container { max-width: 1400px; margin: 0 auto; }
            .header {
                background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                color: white;
                padding: 30px;
                border-radius: 15px;
                margin-bottom: 30px;
            }
            .card {
                background: white;
                padding: 30px;
                border-radius: 12px;
                box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            }
            .back-btn {
                display: inline-block;
                background: white;
                color: #667eea;
                padding: 10px 20px;
                border-radius: 8px;
                text-decoration: none;
                font-weight: 600;
            }
            .back-btn:hover { background: #f0f0f0; }
            ul { line-height: 2; }
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <a href="/" class="back-btn">← Back to Home</a>
                <h1 style="margin-top: 20px;">💳 Collections Tracking</h1>
                <p>Payment monitoring and efficiency metrics</p>
            </div>

            <div class="card">
                <h2>Collections Features in Full System:</h2>
                <ul>
                    <li>✓ Automatic collection record creation from invoices</li>
                    <li>✓ Payment status tracking (Pending, Paid, Overdue)</li>
                    <li>✓ Days outstanding calculation</li>
                    <li>✓ Collection efficiency percentage</li>
                    <li>✓ Balance outstanding tracking</li>
                    <li>✓ Due date monitoring with automatic status updates</li>
                    <li>✓ Linked to customer and salesperson for full context</li>
                </ul>

                <h3 style="margin-top: 30px;">Sample SQL Queries:</h3>
                <pre style="background: #f5f5f5; padding: 15px; border-radius: 8px; overflow-x: auto;">
-- All overdue collections
SELECT * FROM vw_collections_summary
WHERE payment_status = 'Overdue'
ORDER BY days_outstanding DESC;

-- Collection efficiency by salesperson
SELECT owner_name,
       AVG(collection_efficiency) as avg_efficiency,
       SUM(balance_outstanding) as total_outstanding
FROM vw_collections_summary
GROUP BY owner_name
ORDER BY avg_efficiency DESC;
                </pre>
            </div>
        </div>
    </body>
    </html>
    """

# API endpoints for JSON data
@app.get("/api/data/customers")
def api_customers():
    """Get customers as JSON"""
    try:
        wb = load_excel_data()
        ws = wb['Customer Master']
        customers = []

        for row in ws.iter_rows(min_row=5, max_row=30, values_only=True):
            if not row[0]:
                break
            customers.append({
                "customer_id": safe_str(row[0]),
                "client_name": safe_str(row[1]),
                "industry": safe_str(row[2]),
                "type": safe_str(row[3]),
                "status": safe_str(row[6]),
                "hunter_id": safe_str(row[8]),
                "farmer_id": safe_str(row[9])
            })

        return JSONResponse(content={"customers": customers, "count": len(customers)})
    except Exception as e:
        return JSONResponse(content={"error": str(e)}, status_code=500)

if __name__ == "__main__":
    print("=" * 60)
    print("Sales Data Management Hub - Demo Viewer")
    print("=" * 60)
    print("")
    print("Starting server...")
    print("Open your browser and visit: http://localhost:8000")
    print("")
    print("Press Ctrl+C to stop the server")
    print("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=8000, log_level="info")
